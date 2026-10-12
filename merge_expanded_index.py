#!/usr/bin/env python3
"""Consolidated original-web-source index over independently hash-verified media artifact parts."""
import csv, json, hashlib, pathlib, re, collections, zipfile
from urllib.parse import urlparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

root=pathlib.Path("artifacts_downloaded")
dest=pathlib.Path("master_index")
dest.mkdir(exist_ok=True)
catnames={
 "enterprise_recruit":"矿业企业招聘",
 "geology_institution_recruit":"地勘单位招聘",
 "graduate_employment_policy":"毕业生就业政策",
 "geological_industry_policy_data":"矿产行业政策"
}
assert root.exists()
parts=[p for p in root.iterdir() if p.is_dir() and p.name.startswith("expanded-")]
assert len(parts)>=16, (len(parts),[p.name for p in parts])
partmap={}
manifest_by_category={}
failures_by_category={}
verification=[]
for part in sorted(parts):
    actual=[p for p in part.iterdir() if p.is_file()]
    for p in actual:
        if p.name in ("source_manifest.csv","failed_sources.csv","capture_run_info.json"):continue
        if p.name in partmap:
            raise ValueError("Duplicate filename across ZIP parts: "+p.name)
        sha=hashlib.sha256(p.read_bytes()).hexdigest()
        partmap[p.name]=(part.name,p,sha)
    m=part/"source_manifest.csv"
    if m.exists():
        with m.open(encoding="utf-8-sig",newline="") as f:
            records=list(csv.DictReader(f))
        cat=part.name.rsplit("-part",1)[0].removeprefix("expanded-")
        if cat not in manifest_by_category:manifest_by_category[cat]=records
    f=part/"failed_sources.csv"
    if f.exists():
        cat=part.name.rsplit("-part",1)[0].removeprefix("expanded-")
        if cat not in failures_by_category:
            with f.open(encoding="utf-8-sig",newline="") as h: failures_by_category[cat]=list(csv.DictReader(h))

data=[]
failed=[]
all_counts={}
headers=["原始资料类别","公告或政策标题","来源网站","文件类型","所在压缩包","文件名","文件字节数","原网页URL","原附件URL","SHA256","源文件SHA256校验","发布年份提示"]
for cat,records in manifest_by_category.items():
    valid_urls=set()
    for row in records:
        name=pathlib.Path(row["file"]).name
        if name not in partmap:raise ValueError("Missing original file from all parts: "+name)
        part,actual,sha=partmap[name]
        if sha!=row["sha256"]:raise ValueError("Image tampered: "+name)
        if actual.stat().st_size!=int(row["bytes"]):raise ValueError("Incorrect file size: "+name)
        title=row["page_name"]
        yr=re.search(r"20(?:2[5-9]|3\d)",title+" "+name)
        record=[
            catnames.get(cat,cat),title,urlparse(row["source_url"]).netloc,row["asset_type"],
            part.replace("expanded-enterprise_recruit","Geo_Enterprise_Recruit_").replace("expanded-geology_institution_recruit","Geo_Institution_").replace("expanded-graduate_employment_policy","Geo_Employment_Policy_Part_").replace("expanded-geological_industry_policy_data","Geo_IndustryPolicy_"),
            name,int(row["bytes"]),row["source_url"],row["asset_url"],sha,"已通过",yr.group(0) if yr else ""
        ]
        # Need exact downloaded filenames; dynamically compute names from part number with logical mapping
        suffix=part.split("-part")[-1]
        if cat=="enterprise_recruit":record[4]=f"Geo_Enterprise_Recruit_{suffix}.zip"
        elif cat=="geology_institution_recruit":record[4]=f"Geo_Institution_{suffix}.zip"
        elif cat=="graduate_employment_policy":record[4]=f"Geo_Employment_Policy_Part_{suffix}.zip"
        elif cat=="geological_industry_policy_data":record[4]=f"Geo_IndustryPolicy_{suffix}.zip"
        data.append(record)
        if row["asset_type"] in ("ORIGINAL_SITE_VIEWPORT","ORIGINAL_ATTACHMENT"):
            valid_urls.add(row["source_url"])
    counts=collections.Counter(v["asset_type"] for v in records)
    all_counts[cat]={"files":len(records),"actual_sources":len(valid_urls),"by_asset_type":dict(counts),"bytes":sum(int(a["bytes"]) for a in records)}
    failed.extend((catnames.get(cat,cat),f["name"],f["url"],f["failure"],f.get("status",""))
                   for f in failures_by_category.get(cat,[]))

assert len(data)==sum(c["files"] for c in all_counts.values())
assert len(data)==279,len(data)
print("CHECKED_SOURCE_FILES",len(data),"CATEGORIES",len(all_counts),
      "SOURCE_URLS",sum(x["actual_sources"] for x in all_counts.values()),"PARTS",len(parts))
print(json.dumps(all_counts,ensure_ascii=False,indent=2))

with (dest/"全部279份原始资料索引.csv").open("w",encoding="utf-8-sig",newline="") as f:
    w=csv.writer(f)
    w.writerow(headers)
    w.writerows(data)
with (dest/"失败来源记录.csv").open("w",encoding="utf-8-sig",newline="") as f:
    w=csv.writer(f);w.writerow(["资料类别","标题","原网址","失败原因","状态"]);w.writerows(failed)
with (dest/"分类采集结果.json").open("w",encoding="utf-8") as f:
    json.dump(all_counts,f,ensure_ascii=False,indent=2)

wb=Workbook()
summary=wb.active
summary.title="分类总览"
summary.append(["素材类别","网页或附件来源数","原始文件数","总大小 (MB)","原始网站首屏截图","原网站整页截图","原始滚动截图","原网页图片","原始附件","PDF原页渲染"])
for cat,stat in all_counts.items():
    c=stat["by_asset_type"]
    summary.append([catnames[cat],stat["actual_sources"],stat["files"],round(stat["bytes"]/1048576,2),
        c.get("ORIGINAL_SITE_VIEWPORT",0),c.get("ORIGINAL_SITE_FULL",0),c.get("ORIGINAL_SITE_SCROLL",0),
        c.get("ORIGINAL_WEBSITE_IMAGE",0),c.get("ORIGINAL_ATTACHMENT",0),c.get("ORIGINAL_PDF_PAGE_RENDER",0)])
summary.append(["合计",sum(x["actual_sources"] for x in all_counts.values()),len(data),
                round(sum(x["bytes"] for x in all_counts.values())/1048576,2),
                *[sum(x["by_asset_type"].get(k,0) for x in all_counts.values()) for k in
                  ("ORIGINAL_SITE_VIEWPORT","ORIGINAL_SITE_FULL","ORIGINAL_SITE_SCROLL","ORIGINAL_WEBSITE_IMAGE","ORIGINAL_ATTACHMENT","ORIGINAL_PDF_PAGE_RENDER")]])
sheets={
    "enterprise_recruit":"企业招聘原件",
    "geology_institution_recruit":"地勘事业单位招聘",
    "graduate_employment_policy":"毕业生就业政策",
    "geological_industry_policy_data":"矿产行业政策数据"
}
for cat,wsname in sheets.items():
    sh=wb.create_sheet(wsname)
    sh.append(headers)
    for r in data:
        if r[0]!=catnames[cat]:continue
        sh.append(r)
        sh.cell(sh.max_row,8).hyperlink=r[7]
        sh.cell(sh.max_row,8).font=Font(color="006A4E",underline="single")
    sh.auto_filter.ref=sh.dimensions
    sh.freeze_panes="D2"
err=wb.create_sheet("失败来源记录")
err.append(["资料类别","标题","原网址","失败原因","状态"])
for r in failed:err.append(r)
err.auto_filter.ref=err.dimensions
err.freeze_panes="C2"
for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines=False
    sheet.freeze_panes=sheet.freeze_panes or "B2"
    sheet.row_dimensions[1].height=31
    for c in sheet[1]:
        c.fill=PatternFill("solid",fgColor="18334A")
        c.font=Font(name="微软雅黑",bold=True,color="FFFFFF",size=11)
        c.alignment=Alignment(vertical="center",wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if not cell.hyperlink:
                cell.font=Font(name="微软雅黑",size=10,color="1E2832")
            cell.alignment=Alignment(vertical="center",wrap_text=True)
    widths={"分类总览":[26,19,16,17,19,19,18,19,14,19],
            "失败来源记录":[26,50,66,70,20]}
    ws=widths.get(sheet.title,[21,58,26,28,40,66,18,80,80,69,20,16])
    for j,width in enumerate(ws,1):
        sheet.column_dimensions[get_column_letter(j)].width=width
    if sheet.title!="分类总览":
        sheet.auto_filter.ref=sheet.dimensions

out=dest/"资源勘查工程_2026至2027招聘和政策_原始资料总索引.xlsx"
wb.save(out)
check=load_workbook(out,read_only=True)
assert check["分类总览"]["C6"].value==279,(check.sheetnames,check["分类总览"]["C6"].value)
check.close()
with zipfile.ZipFile(dest/"资源勘查工程_新增原始资料总索引包.zip","w",zipfile.ZIP_DEFLATED) as z:
    for f in dest.iterdir():
        if f.is_file() and f.suffix!=".zip":z.write(f,f.name)
print("ALL DONE",str(out),out.stat().st_size)
